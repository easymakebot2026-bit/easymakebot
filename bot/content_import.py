"""Spreadsheet helpers for the "Content List" tool (bot/handlers/tools/content_list.py):
a downloadable sample template, and a parser that maps whatever header row a
user's own file has onto our content fields — including an optional
"Category" column for nested grouping (see bot/db/models.py: ContentItem.parent_id)
and an optional "Code" column that makes re-uploading the same file an
upsert: a row whose Code matches an existing item (in this bot) updates it
in place instead of creating a duplicate — see bot/content_nav.py:upsert_item.
An optional "Action" column with the value "Delete"/"حذف" removes the
matching item instead; deletion never happens implicitly (e.g. a blank
title never means delete), only via this explicit marker.

The parser is deliberately forgiving so a bot owner can upload their *own*
existing catalogue file, not just our sample:
- both .xlsx and .csv are accepted (CSV delimiter is sniffed: comma / semicolon / tab);
- the header row doesn't have to be the first row — the first ~15 rows are
  scanned and the one matching the most known column names wins (lets people
  keep a title/logo/notes block above their table);
- column names are matched loosely and in 10 languages (FA, EN, AR, TR, RU,
  FR, DE, ES, IT, KO — the same set generate_sample_excel offers): Persian/
  Arabic letter variants (ي/ی, ك/ک) are unified, separators and punctuation
  are ignored, and a header that merely *contains* a known name (e.g.
  "Product Title *") still maps.
Nothing here needs a network call or any external service.
"""

import csv
import io
import re
import unicodedata
from io import BytesIO

from openpyxl import Workbook, load_workbook

# Column header synonyms in every language we ship a sample for (see
# _SAMPLE_LANGS), matched after normalization against a file's header row so a
# bot owner's own catalogue works too — not just our generated sample.
_COLUMN_SYNONYMS = {
    "title": {
        # en / generic
        "title", "name", "product", "product name", "item", "item name", "heading", "subject",
        # fa
        "عنوان", "نام", "نام محصول", "کالا", "محصول", "خدمت",
        # ar
        "العنوان", "الاسم", "اسم",
        # tr
        "başlık", "ad", "isim", "ürün adı",
        # ru
        "заголовок", "название", "имя", "наименование",
        # fr
        "titre", "nom", "nom du produit",
        # de
        "titel", "bezeichnung", "produktname",
        # es
        "título", "nombre", "nombre del producto",
        # it
        "titolo", "nome", "nome prodotto",
        # ko
        "제목", "이름", "상품명", "상품", "품목",
    },
    "body": {
        "body", "text", "description", "desc", "descr", "content", "details", "detail",
        "about", "info",
        "متن", "توضیحات", "توضیح", "محتوا", "شرح", "جزئیات",
        "النص", "الوصف", "وصف", "المحتوى", "التفاصيل",
        "metin", "açıklama", "aciklama", "içerik", "icerik", "detay",
        "текст", "описание", "содержание", "подробности",
        "texte", "contenu", "détails", "descriptif",
        "beschreibung", "inhalt", "beschreibungstext",
        "texto", "descripción", "descripcion", "contenido", "detalles",
        "testo", "descrizione", "contenuto", "dettagli",
        "내용", "본문", "설명", "상세", "상세설명",
    },
    "image_url": {
        "image", "image url", "imageurl", "img", "photo", "picture", "thumbnail",
        "لینک تصویر", "عکس", "تصویر", "عکس محصول", "آدرس تصویر",
        "رابط الصورة", "الصورة", "صورة",
        "görsel", "görsel url", "gorsel", "resim", "fotoğraf", "fotograf", "foto",
        "изображение", "ссылка на изображение", "картинка", "фото", "фотография",
        "url de l'image", "url image", "illustration",
        "bild", "bild-url", "bild url", "bildurl", "abbildung",
        "imagen", "url de imagen", "url imagen", "fotografía", "fotografia",
        "immagine", "url immagine",
        "이미지", "이미지 url", "사진", "그림",
    },
    "link_url": {
        "link", "url", "website", "web site", "web", "page", "href", "weblink",
        "لینک", "آدرس", "وبسایت", "وب سایت", "لینک صفحه",
        "الرابط", "رابط", "الموقع", "الصفحة",
        "bağlantı", "baglanti", "web sitesi", "adres", "sayfa",
        "ссылка", "сайт", "веб-сайт", "вебсайт", "страница",
        "lien", "site", "site web", "page web",
        "webseite", "seite", "adresse", "webadresse",
        "enlace", "sitio web", "página", "pagina", "vínculo", "vinculo",
        "sito web", "collegamento",
        "링크", "웹사이트", "주소", "페이지",
    },
    "category": {
        "category", "categories", "parent", "parent category", "group", "section", "folder",
        "دسته", "دسته‌بندی", "دسته بندی", "دسته‌ی والد", "گروه", "بخش", "والد", "زیرگروه",
        "التصنيف", "الفئة", "القسم", "المجموعة", "تصنيف",
        "kategori", "grup", "bölüm", "bolum", "üst", "üst kategori", "ust kategori",
        "категория", "группа", "раздел", "родитель", "родительская категория",
        "catégorie", "categorie", "groupe", "rubrique",
        "kategorie", "gruppe", "bereich", "ordner", "übergeordnet", "uebergeordnet",
        "oberkategorie",
        "categoría", "categoria", "grupo", "sección", "seccion", "padre",
        "gruppo", "sezione", "genitore",
        "분류", "카테고리", "그룹", "섹션", "상위", "상위 분류",
    },
    "code": {
        "code", "sku", "id", "ref", "reference", "shortcut", "short code",
        "کد", "شناسه", "کد کوتاه",
        "الرمز", "الكود", "رمز", "المعرف",
        "kod", "kimlik", "kısa kod",
        "код", "артикул", "идентификатор", "ид",
        "référence", "identifiant",
        "kürzel", "kuerzel", "artikelnummer", "kennung",
        "código", "codigo", "referencia", "identificador",
        "codice", "riferimento",
        "코드", "식별자", "참조",
    },
    "action": {
        "action", "op", "operation", "operations", "command", "cmd",
        "عملیات", "اقدام", "دستور",
        "الإجراء", "الأمر", "العملية", "إجراء",
        "işlem", "islem", "eylem", "komut",
        "действие", "операция", "команда",
        "opération",
        "aktion", "vorgang", "befehl", "aktionen",
        "acción", "accion", "operación", "operacion", "comando", "acciones",
        "azione", "operazione",
        "작업", "동작", "명령", "작업내용",
    },
}

# "Action" cell values that mean "delete the item matching this Code".
_DELETE_ACTION_VALUES = {
    "delete", "remove", "del", "drop", "erase",
    "حذف", "پاک", "پاک کردن", "حذف کردن",
    "إزالة", "احذف", "امسح",
    "sil", "kaldır", "kaldir", "çıkar", "cikar",
    "удалить", "удаление", "убрать",
    "supprimer", "effacer", "retirer", "suppr",
    "löschen", "loeschen", "entfernen", "lösche",
    "eliminar", "borrar", "quitar",
    "elimina", "eliminare", "rimuovi", "cancella",
    "삭제", "제거", "지우기",
}

# One localized sample per language. Each has the seven headers in that
# language and the same worked example: a two-level nesting (top category >
# sub-category > one real item with a Code, so re-uploading updates it in
# place instead of duplicating). `rtl` flips the sheet for FA/AR.
_SAMPLE_LANGS: dict[str, dict] = {
    "fa": {
        "label": "🇮🇷 فارسی",
        "rtl": True,
        "headers": ["عنوان", "متن", "لینک تصویر", "لینک", "دسته‌بندی", "کد", "عملیات"],
        "example": ["ورزش", "فوتبال", "نمونه آیتم", "این متنی است که به کاربران نمایش داده می‌شود."],
    },
    "en": {
        "label": "🇬🇧 English",
        "rtl": False,
        "headers": ["Title", "Body", "Image URL", "Link", "Category", "Code", "Action"],
        "example": ["Sports", "Football", "Example item", "This is the body text shown to users."],
    },
    "ar": {
        "label": "🇸🇦 العربية",
        "rtl": True,
        "headers": ["العنوان", "النص", "رابط الصورة", "الرابط", "التصنيف", "الرمز", "الإجراء"],
        "example": ["الرياضة", "كرة القدم", "عنصر مثال", "هذا هو النص الذي يظهر للمستخدمين."],
    },
    "tr": {
        "label": "🇹🇷 Türkçe",
        "rtl": False,
        "headers": ["Başlık", "Metin", "Görsel URL", "Bağlantı", "Kategori", "Kod", "İşlem"],
        "example": ["Spor", "Futbol", "Örnek öğe", "Bu, kullanıcılara gösterilen metindir."],
    },
    "ru": {
        "label": "🇷🇺 Русский",
        "rtl": False,
        "headers": ["Заголовок", "Текст", "Ссылка на изображение", "Ссылка", "Категория", "Код", "Действие"],
        "example": ["Спорт", "Футбол", "Пример элемента", "Это текст, который видят пользователи."],
    },
    "fr": {
        "label": "🇫🇷 Français",
        "rtl": False,
        "headers": ["Titre", "Texte", "URL de l'image", "Lien", "Catégorie", "Code", "Action"],
        "example": ["Sport", "Football", "Exemple d'élément", "Ceci est le texte affiché aux utilisateurs."],
    },
    "de": {
        "label": "🇩🇪 Deutsch",
        "rtl": False,
        "headers": ["Titel", "Text", "Bild-URL", "Link", "Kategorie", "Code", "Aktion"],
        "example": ["Sport", "Fußball", "Beispieleintrag", "Dies ist der Text, der Benutzern angezeigt wird."],
    },
    "es": {
        "label": "🇪🇸 Español",
        "rtl": False,
        "headers": ["Título", "Texto", "URL de imagen", "Enlace", "Categoría", "Código", "Acción"],
        "example": ["Deportes", "Fútbol", "Elemento de ejemplo", "Este es el texto que se muestra a los usuarios."],
    },
    "it": {
        "label": "🇮🇹 Italiano",
        "rtl": False,
        "headers": ["Titolo", "Testo", "URL immagine", "Link", "Categoria", "Codice", "Azione"],
        "example": ["Sport", "Calcio", "Elemento di esempio", "Questo è il testo mostrato agli utenti."],
    },
    "ko": {
        "label": "🇰🇷 한국어",
        "rtl": False,
        "headers": ["제목", "내용", "이미지 URL", "링크", "분류", "코드", "작업"],
        "example": ["스포츠", "축구", "예시 항목", "사용자에게 표시되는 본문 텍스트입니다."],
    },
}

DEFAULT_SAMPLE_LANG = "en"

# [(code, "🇬🇧 English"), ...] in a stable display order — for the language
# picker keyboard (bot/keyboards.py: content_sample_lang_keyboard).
SAMPLE_LANGUAGES = [(code, spec["label"]) for code, spec in _SAMPLE_LANGS.items()]

# Back-compat: some callers/tests import these directly.
SAMPLE_HEADERS = _SAMPLE_LANGS[DEFAULT_SAMPLE_LANG]["headers"]

# How many leading rows to scan when hunting for the header row.
_HEADER_SCAN_ROWS = 15


def _sample_rows(lang: str) -> list[list[str]]:
    top, sub, item, body = _SAMPLE_LANGS[lang]["example"]
    return [
        [top, "", "", "", "", "", ""],
        [sub, "", "", "", top, "", ""],
        [item, body, "https://example.com/image.jpg", "https://example.com", sub, "101", ""],
    ]


def generate_sample_excel(lang: str = DEFAULT_SAMPLE_LANG) -> bytes:
    """A ready-to-fill .xlsx whose column headers are in `lang` (one of
    _SAMPLE_LANGS; unknown codes fall back to English)."""
    if lang not in _SAMPLE_LANGS:
        lang = DEFAULT_SAMPLE_LANG
    spec = _SAMPLE_LANGS[lang]

    wb = Workbook()
    ws = wb.active
    ws.title = "Content"
    ws.sheet_view.rightToLeft = spec["rtl"]
    ws.append(spec["headers"])
    for row in _sample_rows(lang):
        ws.append(row)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _normalize(value: object) -> str:
    """Fold a header cell to a comparable form: NFKC, lowercased, trimmed,
    Persian/Arabic yeh & kaf unified, and runs of separators/punctuation
    collapsed to a single space."""
    text = unicodedata.normalize("NFKC", str(value)).strip().lower()
    text = text.replace("ي", "ی").replace("ك", "ک").replace("‌", " ")
    text = re.sub(r"[\s_\-./\\|()\[\]{}:*#\"']+", " ", text)
    return text.strip()


_NORMALIZED_SYNONYMS = {
    field: {_normalize(name) for name in names} for field, names in _COLUMN_SYNONYMS.items()
}


def _match_field(header: object) -> str | None:
    if header is None:
        return None
    norm = _normalize(header)
    if not norm:
        return None
    for field, names in _NORMALIZED_SYNONYMS.items():
        if norm in names:
            return field
    # Loose fallback: a header that contains a known name, or is contained
    # by one. Both sides must be >= 3 chars so stray single-letter data
    # cells ("a", "b") can't accidentally match "name"/"body" and get
    # mistaken for a header row.
    if len(norm) >= 3:
        for field, names in _NORMALIZED_SYNONYMS.items():
            for name in names:
                if len(name) >= 3 and (name in norm or norm in name):
                    return field
    return None


def _pick_header_row(rows: list[tuple]) -> tuple[int, list[str | None]]:
    """Among the first _HEADER_SCAN_ROWS rows, return (index, column_fields)
    for the row that maps to the most known fields. Ties go to the earliest
    row. Raises ValueError if no row maps to anything."""
    best_index = -1
    best_fields: list[str | None] = []
    best_score = 0
    for index, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
        fields = [_match_field(cell) for cell in row]
        score = sum(1 for f in fields if f)
        if score > best_score:
            best_index, best_fields, best_score = index, fields, score

    if best_score == 0:
        raise ValueError(
            "None of the column headers were recognized. Use the sample "
            "template's headers (Title, Body, Image URL, Link, Category, Code, Action) "
            "or matching names, in either a .xlsx or .csv file."
        )
    return best_index, best_fields


def _looks_like_csv(data: bytes) -> bool:
    # .xlsx is a zip ("PK\x03\x04"); .xls is an OLE2 doc ("\xD0\xCF\x11\xE0").
    # Anything else that decodes as text and has a comma/semicolon/tab in its
    # first line we treat as CSV.
    if data[:2] == b"PK" or data[:4] == b"\xd0\xcf\x11\xe0":
        return False
    try:
        head = data[:4096].decode("utf-8-sig")
    except UnicodeDecodeError:
        return False
    first_line = head.splitlines()[0] if head.splitlines() else ""
    return any(sep in first_line for sep in (",", ";", "\t"))


def _load_rows(data: bytes, filename: str | None) -> list[tuple]:
    """Return every non-empty row of the file as a tuple, from either an
    .xlsx workbook (first/active sheet) or a CSV (delimiter sniffed)."""
    name = (filename or "").lower()
    is_csv = name.endswith(".csv") or (
        not name.endswith((".xlsx", ".xlsm", ".xltx", ".xls")) and _looks_like_csv(data)
    )

    if is_csv:
        text = data.decode("utf-8-sig", errors="replace")
        sample = text[:4096]
        try:
            dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        rows = [tuple(cell.strip() for cell in row) for row in reader]
    else:
        try:
            wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise ValueError(
                "Could not read this file — send it as an .xlsx or .csv spreadsheet."
            ) from exc
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))

    non_empty = [
        row for row in rows if row and any(cell is not None and str(cell).strip() for cell in row)
    ]
    if not non_empty:
        raise ValueError("The spreadsheet is empty.")
    return non_empty


def parse_content_excel(data: bytes, filename: str | None = None) -> list[dict]:
    """Returns a list of {"title", "body", "image_url", "link_url", "category",
    "code", "action"} dicts, one per data row — "category" is the *title
    text* of another row in this same file (its intended parent), or None
    for a top-level item; the caller resolves it to an actual parent_id.
    "action" is "delete" or None (see module docstring). Rows without a
    title are skipped, EXCEPT delete rows, which only need a Code.

    Accepts .xlsx or .csv; the header row need not be the first row.
    Raises ValueError if the file can't be read or has no recognizable
    header row. `filename` (when known) only steers the .xlsx-vs-.csv guess.
    """
    rows = _load_rows(data, filename)
    header_index, column_fields = _pick_header_row(rows)
    data_rows = rows[header_index + 1 :]

    items = []
    for row in data_rows:
        item = {
            "title": None, "body": None, "image_url": None, "link_url": None,
            "category": None, "code": None, "action": None,
        }
        for field, value in zip(column_fields, row):
            if field is None or value is None:
                continue
            text = str(value).strip()
            if text:
                item[field] = text

        item["action"] = (
            "delete" if (item["action"] or "").strip().lower() in _DELETE_ACTION_VALUES else None
        )

        if item["action"] == "delete":
            if not item["code"]:
                continue  # a delete row needs a Code to know what to delete
        elif not item["title"]:
            continue

        item["body"] = item["body"] or ""
        item["category"] = item["category"] or None
        item["code"] = item["code"] or None
        items.append(item)

    return items
